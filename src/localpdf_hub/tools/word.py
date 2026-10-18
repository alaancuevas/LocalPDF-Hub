"""PDF to Word, and Word to PDF through whatever converter the system has."""

import io
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Inches

from localpdf_hub import errors
from localpdf_hub.errors import AppError
from localpdf_hub.limits import check_deadline
from localpdf_hub.platform import DocumentConverter
from localpdf_hub.tools.pdf_files import read_pdf
from localpdf_hub.workspace import without_prefix

_RENDER_SCALE = 1.5
_IMAGE_WIDTH_IN = 6.0


def to_word(path: Path, title: str, out_dir: Path) -> Path:
    """Copies the text page by page; scanned PDFs without text become page images."""
    reader = read_pdf(path)
    texts = []
    for page in reader.pages:
        check_deadline()
        texts.append((page.extract_text() or "").strip())

    document = Document()
    document.add_heading(title, level=1)
    if any(texts):
        _add_text_pages(document, texts)
    else:
        _add_page_images(document, path)
    out = out_dir / f"{title}.docx"
    document.save(str(out))
    return out


def _add_text_pages(document: DocxDocument, texts: list[str]) -> None:
    for index, text in enumerate(texts):
        if text:
            document.add_paragraph(text)
        if index < len(texts) - 1:
            document.add_page_break()


def _add_page_images(document: DocxDocument, path: Path) -> None:
    pdf = pdfium.PdfDocument(path)
    try:
        for index in range(len(pdf)):
            check_deadline()
            page = pdf[index]
            try:
                png = io.BytesIO()
                page.render(scale=_RENDER_SCALE).to_pil().save(png, "PNG")
            finally:
                page.close()
            png.seek(0)
            document.add_picture(png, width=Inches(_IMAGE_WIDTH_IN))
            if index < len(pdf) - 1:
                document.add_page_break()
    finally:
        pdf.close()


def from_word(documents: list[Path], converter: DocumentConverter | None, out_dir: Path) -> Path:
    """Returns the PDF itself for one document, or a ZIP for several."""
    if converter is None:
        raise AppError(
            errors.CONVERTER_MISSING,
            "Word to PDF needs Microsoft Word (Windows) or LibreOffice installed.",
        )
    pdfs = converter.convert(documents, out_dir)
    if len(pdfs) == 1:
        return pdfs[0]
    out = out_dir / "converted.zip"
    used: set[str] = set()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for pdf in pdfs:
            # Uploads are stored as "000_name.docx"; the user expects "name.pdf",
            # unless two uploads had the same name.
            name = without_prefix(pdf.name)
            if name in used:
                name = pdf.name
            used.add(name)
            archive.write(pdf, arcname=name)
    return out
